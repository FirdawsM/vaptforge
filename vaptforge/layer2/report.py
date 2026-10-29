import json
from datetime import datetime
from pathlib import Path

from docx import Document

from vaptforge.layer1.scoring import command_failed
from .report_sections import (
    add_attack_surface,
    add_findings,
    add_not_run_appendix,
    add_summary,
    sanitize_for_docx,
)


DEFAULT_TEMPLATE = Path(__file__).resolve().parent / "templates" / "jamaa_template.docx"


def _cover_date(scanned_at):
    try:
        value = datetime.fromisoformat(str(scanned_at).replace("Z", "+00:00"))
    except (TypeError, ValueError):
        return str(scanned_at or "Scan Date")
    day = value.day
    suffix = "th" if 10 <= day % 100 <= 20 else {1: "st", 2: "nd", 3: "rd"}.get(day % 10, "th")
    return day, suffix, value.strftime(" %B, %Y")


def _set_cover_date(document, scanned_at):
    paragraph = document.paragraphs[3]
    runs = paragraph.runs
    date = _cover_date(scanned_at)
    if isinstance(date, str):
        runs[1].text = date
        runs[2].text = ""
        runs[3].text = ""
        runs[4].text = ""
    else:
        day, suffix, rest = date
        runs[1].text = str(day)
        runs[2].text = suffix
        runs[3].text = " "
        runs[4].text = rest.lstrip()


def _set_prepared_by(document, prepared_by):
    if prepared_by is None or not str(prepared_by).strip():
        return
    paragraphs = document.paragraphs
    label_index = next(
        (index for index, paragraph in enumerate(paragraphs) if "Prepared by" in paragraph.text),
        None,
    )
    if label_index is None or label_index + 1 >= len(paragraphs):
        raise ValueError("The report template does not contain a Prepared by field")
    runs = paragraphs[label_index + 1].runs
    target_run = next((run for run in reversed(runs) if run.text.strip()), None)
    if target_run is None:
        raise ValueError("The report template has an empty Prepared by field")
    target_run.text = sanitize_for_docx(str(prepared_by).strip())


def build_report(results_path, template_path=None, output_path="report.docx", prepared_by=None):
    results_path = Path(results_path)
    template_path = Path(template_path) if template_path else DEFAULT_TEMPLATE
    output_path = Path(output_path)
    if not results_path.is_file():
        raise FileNotFoundError(f"Scan results file not found: {results_path}")
    if not template_path.is_file():
        raise FileNotFoundError(
            f"Jamaa Word template not found: {template_path}. Provide the approved template with --template."
        )

    data = json.loads(results_path.read_text(encoding="utf-8"))
    if not isinstance(data, dict) or not isinstance(data.get("results", []), list):
        raise ValueError("scan_results.json must be an object containing a results array")
    for result in data.get("results", []):
        if result.get("result") in ("Pass", "Fail") and command_failed(
            result.get("raw_output") or "", result.get("exit_code")
        ):
            result["result"] = "Error"
            result["note"] = "Captured output or exit code indicates the command did not complete successfully; no verdict is possible."

    document = Document(str(template_path))
    _set_cover_date(document, data.get("scanned_at"))
    _set_prepared_by(document, prepared_by)
    add_summary(document, data)
    add_attack_surface(document, data)
    add_findings(document, data)
    add_not_run_appendix(document, data)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))
    return output_path
