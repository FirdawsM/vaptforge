from collections import Counter, defaultdict
import re

from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


PURPLE = "7030A0"
WHITE = "FFFFFF"
RESULT_STYLES = {
    "Pass": ("C6EFCE", "006100"),
    "Fail": ("FFC7CE", "9C0006"),
    "Manual Review": ("FFEB9C", "9C6500"),
    "Error": ("FFEB9C", "9C6500"),
}
PHASE_LABELS = ("0-2 dynamic scan", "3 vuln validation", "4 HIGH RISK")
NOT_RUN_PREFIXES = ("Skipped (", "Blocked (", "Error")
_ANSI_ESCAPE = re.compile(r"\x1b(?:\[[0-?]*[ -/]*[@-~]|\][^\x07]*(?:\x07|\x1b\\))")
_XML_UNSAFE = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F\x7F-\x9F\uD800-\uDFFF\uFFFE\uFFFF]")


def sanitize_for_docx(text):
    text = str(text)
    text = _ANSI_ESCAPE.sub("", text)
    return _XML_UNSAFE.sub("", text)


def _phase_names(max_phase):
    try:
        ceiling = int(max_phase)
    except (TypeError, ValueError):
        return "Unknown"
    included = []
    if ceiling >= 0:
        included.append(PHASE_LABELS[0])
    if ceiling >= 3:
        included.append(PHASE_LABELS[1])
    if ceiling >= 4:
        included.append(PHASE_LABELS[2])
    return ", ".join(included) if included else "None"


def _display(value, fallback="Not provided"):
    return sanitize_for_docx(value) if value not in (None, "") else fallback


def _set_cell(cell, value, header=False):
    cell.text = _display(value)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.space_before = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.5
        for run in paragraph.runs:
            run.font.size = Pt(10)
            if header:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string(WHITE)
    if header:
        properties = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), PURPLE)
        properties.append(shading)


def _new_table(document, headers, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = False
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    header_row = table.rows[0]
    header_row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for index, (cell, value) in enumerate(zip(header_row.cells, headers)):
        cell.width = Inches(widths[index])
        _set_cell(cell, value, header=True)
    return table


def _add_row(table, values):
    row = table.add_row()
    row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
    for index, (cell, value) in enumerate(zip(row.cells, values)):
        cell.width = table.columns[index].width
        _set_cell(cell, value)


def _body_paragraph(document, text):
    paragraph = document.add_paragraph(sanitize_for_docx(text))
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(5)
    paragraph.paragraph_format.line_spacing = 1.5
    for run in paragraph.runs:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor.from_string("000000")
    return paragraph


def add_summary(document, data):
    document.add_heading("1. Executive Summary", level=1)
    _body_paragraph(document, f"Target: {_display(data.get('target'))}")
    _body_paragraph(document, f"Scanned at: {_display(data.get('scanned_at'))}")
    _body_paragraph(document, f"Phases run: {_phase_names(data.get('max_phase'))}")
    if data.get("signoff_ref"):
        _body_paragraph(document, f"Sign-off reference: {data['signoff_ref']}")

    counts = Counter(result.get("result", "Unknown") for result in data.get("results", []))
    summary = " | ".join(
        f"{label}: {counts.get(label, 0)}" for label in ("Pass", "Fail", "Manual Review", "Error")
    )
    paragraph = _body_paragraph(document, "")
    run = paragraph.add_run(summary)
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string("000000")


def add_attack_surface(document, data):
    document.add_heading("2. Attack Surface Mapping", level=1)
    _body_paragraph(
        document,
        "Coverage describes the test catalog, not this scan's verification of site components."
    )
    coverage = data.get("attack_surface_coverage") or {}
    by_component = coverage.get("by_component") or {}
    table = _new_table(
        document,
        ("Component", "Elements covered", "Elements untested"),
        (1.35, 2.52, 2.52),
    )
    for component, details in by_component.items():
        _add_row(
            table,
            (
                component,
                ", ".join(map(str, details.get("elements_covered") or [])) or "None",
                ", ".join(map(str, details.get("elements_untested") or [])) or "None",
            ),
        )
    untagged = coverage.get("untagged_tests") or []
    if untagged:
        _body_paragraph(document, f"Untagged tests: {', '.join(map(str, untagged))}")


def _add_result(document, result, scanned_at):
    heading = document.add_heading(level=3)
    heading.add_run(
        f"{_display(result.get('wstg_id'), 'Unknown WSTG ID')} - "
        f"{_display(result.get('name'), 'Unnamed test')} - Phase {_display(result.get('phase'), 'Unknown')}"
    )
    status = _body_paragraph(document, f"Result: {_display(result.get('result'), 'Unknown')}")
    status_run = status.runs[0]
    status_run.bold = True
    colors = RESULT_STYLES.get(result.get("result"))
    if colors:
        fill, text = colors
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        status._p.get_or_add_pPr().append(shading)
        status_run.font.color.rgb = RGBColor.from_string(text)
    note = sanitize_for_docx(result.get("note") or "").lower()
    explains_review = any(phrase in note for phrase in ("human review", "manual review", "review raw output"))
    if result.get("result") == "Manual Review" and not explains_review:
        _body_paragraph(document, "Human review required: review the raw output before client delivery.")
    if result.get("note"):
        _body_paragraph(document, f"Note: {result['note']}")
    if result.get("caution"):
        _body_paragraph(document, f"Caution: {result['caution']}")
    if result.get("command"):
        _body_paragraph(document, f"Command run: {result['command']}")
    try:
        phase4 = int(result.get("phase")) == 4
    except (TypeError, ValueError):
        phase4 = False
    if phase4:
        _body_paragraph(document, f"Scanned at: {_display(scanned_at)}")
        _body_paragraph(document, f"Sign-off reference: {_display(result.get('signoff_ref'))}")
    _body_paragraph(document, "Raw output:")
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Pt(12)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.15
    raw_output = result.get("raw_output") or "(No output captured.)"
    run = paragraph.add_run(sanitize_for_docx(raw_output))
    run.font.name = "Courier New"
    run.font.size = Pt(8)


def add_findings(document, data):
    document.add_heading("3. Findings", level=1)
    grouped = defaultdict(list)
    for result in data.get("results", []):
        if str(result.get("result", "")).startswith(NOT_RUN_PREFIXES):
            continue
        grouped[_display(result.get("component"), "Uncategorized")].append(result)
    if not grouped:
        _body_paragraph(document, "No tests were run in this scan.")
    for component in sorted(grouped):
        document.add_heading(component, level=2)
        for result in grouped[component]:
            _add_result(document, result, data.get("scanned_at"))


def add_not_run_appendix(document, data):
    skipped = [
        result for result in data.get("results", [])
        if str(result.get("result", "")).startswith(NOT_RUN_PREFIXES)
    ]
    if not skipped:
        return
    document.add_heading("Appendix: Tests Not Run or Inconclusive", level=1)
    _body_paragraph(
        document,
        "Skipped and blocked tests were not run. Error results were attempted but did not complete successfully. None should be read as passing.",
    )
    for result in skipped:
        text = (
            f"{_display(result.get('wstg_id'), 'Unknown WSTG ID')} - "
            f"{_display(result.get('name'), 'Unnamed test')}. "
            f"{_display(result.get('result'), 'Not run')}. "
            f"Reason: {_display(result.get('note'))}"
        )
        paragraph = _body_paragraph(document, f"• {text}")
        paragraph.paragraph_format.space_before = Pt(5)
        paragraph.paragraph_format.space_after = Pt(5)
        paragraph.paragraph_format.line_spacing = 1.5
        paragraph.paragraph_format.left_indent = Inches(0.25)
        for run in paragraph.runs:
            run.font.size = Pt(12)
        if str(result.get("result", "")).startswith("Error"):
            if result.get("command"):
                _body_paragraph(document, f"Command run: {result['command']}")
            _body_paragraph(document, "Captured output:")
            output = document.add_paragraph()
            output.paragraph_format.left_indent = Pt(12)
            output.paragraph_format.space_after = Pt(8)
            output.paragraph_format.line_spacing = 1.0
            run = output.add_run(sanitize_for_docx(result.get("raw_output") or "(No output captured.)"))
            run.font.name = "Courier New"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor.from_string("000000")
